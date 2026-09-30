from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"D:\Desktop\AI伯乐智能体\整合平台1.0\论文调研进展报告.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.4)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.6)
section.right_margin = Cm(2.6)

styles = doc.styles
styles["Normal"].font.name = "宋体"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(11)

title = styles["Title"]
title.font.name = "黑体"
title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
title.font.size = Pt(22)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
    style = styles[name]
    style.font.name = "黑体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(13)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

body = styles["Body Text"]
body.font.name = "宋体"
body._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
body.font.size = Pt(11)
body.paragraph_format.first_line_indent = Cm(0.74)
body.paragraph_format.line_spacing = 1.45
body.paragraph_format.space_after = Pt(6)
body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

bullet = styles["List Bullet"]
bullet.font.name = "宋体"
bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
bullet.font.size = Pt(10.5)
bullet.paragraph_format.left_indent = Cm(0.74)
bullet.paragraph_format.first_line_indent = Cm(-0.42)
bullet.paragraph_format.line_spacing = 1.3
bullet.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("论文调研进展报告")
p_ppr = p._p.get_or_add_pPr()
p_border = p_ppr.find(qn("w:pBdr"))
if p_border is not None:
    p_ppr.remove(p_border)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("儿童能力画像与儿童 AI 交互方向")
r.font.name = "黑体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(80, 80, 80)
p.paragraph_format.space_after = Pt(18)

add_body(doc, "本周围绕项目可能形成论文的研究空间开展了初步调研。调研先从与项目相关的研究方向入手，梳理代表性论文的研究动机、系统设计、方法与贡献，再将已有工作与当前项目进行比较，重点寻找可以借鉴的研究路径、尚未被充分讨论的问题，以及能够进一步转化为论文研究问题的切入点。目前已形成四类文献主题，并在儿童纠错、家长参与和双端呈现等方向上完成了进一步阅读。")

doc.add_heading("一 调研目标与思路", level=1)
add_body(doc, "本阶段调研的主要目标不是完整覆盖某一领域，而是围绕项目快速建立研究坐标：了解相近系统已经解决了什么问题、采用了什么研究方法、贡献通常如何表述，并据此判断我们的项目还可以补充哪些实证工作或设计贡献。")
add_body(doc, "具体过程分为三步：第一，查找儿童能力评估、儿童与生成式 AI 交互、家庭参与等相关论文；第二，比较这些工作与项目在研究对象、交互流程、证据来源和结果呈现上的共性与差异；第三，根据比较结果提出可能的论文方向，并继续阅读与这些方向直接相关的论文。")

doc.add_heading("二 已调研的主要内容", level=1)
doc.add_heading("1 多元能力与游戏化评估", level=2)
add_body(doc, "该方向主要关注如何把多元智能或兴趣特征转化为可观察、可记录的指标，并通过数字化系统完成评分、可视化和建议生成。已阅读的 MULIA 等工作表明，相关研究通常强调指标设计、专家验证、评分一致性和系统可用性，并谨慎地将结果表述为可变化的学习偏好，而不是对儿童能力作永久诊断。另一类在线识别系统则通过问卷和简单加权形成多维画像，但普遍较依赖自陈数据，缺少活动过程与行为证据的共同解释。")

doc.add_heading("2 儿童与 AI 共创 学习和过程证据", level=2)
add_body(doc, "该方向关注 AI 进入故事创作、阅读、绘画、编程或问题解决活动后，如何影响儿童的参与、表达、学习与创造。已有工作主要把 AI 视为共创伙伴或学习支架，通过追问、提示和任务调整降低参与门槛。与项目相关的启发是，研究重点可以从最终得分转向儿童在自然活动中的行为过程，并将过程证据用于解释画像，而不是只给出一个结论。")

doc.add_heading("3 儿童对 AI 的理解 信任和纠错", level=2)
add_body(doc, "该方向主要讨论儿童如何理解 AI 推理、是否信任 AI 信息，以及在发现错误时如何反应。相关研究显示，儿童会持续追问，也能够在一定条件下识别并质疑 AI 错误，但总体上仍可能对 AI 保持较高信任。关于儿童能动性的综述进一步提出，儿童应能够拒绝建议、修改输出、要求重新生成，并理解系统能力与局限。这为项目中的儿童纠错机制提供了较直接的理论依据。")

doc.add_heading("4 家长参与 双端呈现和家庭协商", level=2)
add_body(doc, "该方向关注 AI 进入家庭后，儿童、家长、教师或专业人员如何共同理解信息并作出决定。已进一步阅读家长需求调研、StoryMate、SET-PAiREd 和 Teddy AI 等工作。它们共同说明，家长不应只是查看后台数据的旁观者，而应成为儿童活动的支持者、信息解释者或共同讲述者；同时，家长参与程度需要根据时间、能力和动机等情境变化，儿童端与家长端也应采用差异化的内容与交互方式。")

doc.add_heading("三 代表性论文与项目关联", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(5.1), Cm(5.0), Cm(5.3)]
headers = ["代表性工作", "主要关注点", "对项目的启发"]
for i, cell in enumerate(table.rows[0].cells):
    cell.width = widths[i]
    cell.text = headers[i]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "2F5597")
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.size = Pt(9.5)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
set_repeat_table_header(table.rows[0])

rows = [
    ("MULIA", "将多元智能转化为幼儿可观察行为，并进行多层验证", "画像需要可解释的行为指标，并明确结果边界"),
    ("Children’s Mental Models of AI Reasoning", "研究儿童如何理解和质疑 AI 推理", "把纠错从识别 AI 错误迁移到质疑 AI 对自己的画像"),
    ("SET-PAiREd", "根据技能 动力和时间设计不同家长参与情境", "家长反问应情境化，重点是校准而不是增加提问数量"),
    ("Teddy AI", "儿童对话界面与家长仪表板的双端架构", "比较儿童端与家长端的差异化呈现及反馈回流"),
]
for ri, data in enumerate(rows):
    cells = table.add_row().cells
    for i, (cell, value) in enumerate(zip(cells, data)):
        cell.width = widths[i]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if ri % 2 == 1:
            set_cell_shading(cell, "EEF3F8")
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            for run in p.runs:
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                run.font.size = Pt(9.3)

doc.add_heading("四 阶段性发现", level=1)
add_bullet(doc, "现有儿童能力评估研究多以问卷、教师观察或固定指标形成结果，项目可以进一步强调活动过程证据，并解释画像结论如何由具体行为产生。")
add_bullet(doc, "对儿童的描述不宜被包装成稳定、永久的能力诊断，更适合表述为当前活动中的行为倾向、兴趣线索或可继续探索的方向。")
add_bullet(doc, "儿童并非只能接受 AI 结论。拒绝、修改、举出反例和决定是否保留证据，都可以成为可观察的能动性指标。")
add_bullet(doc, "家长端不应只是儿童结果的展示后台。更有研究价值的问题是，家长如何补充证据、理解不确定性，并与儿童共同协商画像。")
add_bullet(doc, "同一份数据需要面向儿童和家长采用不同叙事。儿童端应支持理解与表达异议，家长端应呈现证据、解释边界并支持后续行动。")

doc.add_heading("五 初步形成的论文方向", level=1)
add_body(doc, "方向一 过程证据驱动的儿童画像。研究如何从故事共创、阅读或其他自然活动中提取行为证据，并将其组织为可理解、可追溯的倾向描述。重点比较仅展示结论与同时展示行为证据时，儿童和家长对画像的理解与信任差异。", bold_lead="方向一 过程证据驱动的儿童画像。")
add_body(doc, "方向二 儿童对 AI 自我画像的纠错与能动性。把已有研究中儿童识别 AI 推理错误的问题迁移到自我表征场景，考察儿童能否理解画像来源、拒绝不符合自己的描述、给出反例，并使系统结果发生可见变化。", bold_lead="方向二 儿童对 AI 自我画像的纠错与能动性。")
add_body(doc, "方向三 家长作为画像校准者的参与机制。研究家长在 AI 画像与日常观察一致、互补或冲突时的接受度和行动反应，并探索何时提问、问什么、如何让家长补充证据，而不是简单增加家长控制。", bold_lead="方向三 家长作为画像校准者的参与机制。")
add_body(doc, "方向四 双端差异化呈现与家庭协商。围绕同一组儿童行为数据，设计儿童端和家长端的差异化报告，比较家长是否能看到行为证据、是否能反馈修改等变量对信任、理解和亲子讨论的影响。", bold_lead="方向四 双端差异化呈现与家庭协商。")

doc.add_heading("六 当前进展与下一步计划", level=1)
add_body(doc, "目前已完成相关研究方向的初步划分，梳理了多篇代表性论文的研究贡献与共性，并围绕家长参与、儿童纠错和双端呈现进行了进一步阅读。现阶段较有潜力的主线是：以活动过程证据形成儿童画像，同时让儿童拥有纠错权，让家长以校准者和共同讲述者的方式参与。该主线能够把项目现有功能与儿童能动性、多利益相关者参与和可解释画像等研究议题连接起来。")
add_body(doc, "下一步将继续收敛论文问题，优先明确研究对象、核心变量和评价指标；补充与儿童自我表征、家庭协商和可解释 AI 相关的文献；在项目流程中标出可直接用于实验对比的关键环节，并据此形成一版研究问题与实验设计草案。")

footer = section.footer
footer_p = footer.paragraphs[0]
add_page_number(footer_p)
for run in footer_p.runs:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)

doc.core_properties.title = "论文调研进展报告"
doc.core_properties.subject = "儿童能力画像与儿童 AI 交互方向"
doc.save(OUT)
print(OUT)
